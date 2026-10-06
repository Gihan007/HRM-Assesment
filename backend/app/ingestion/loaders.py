"""File loaders for PDF, DOCX, TXT, MD -> extracted text with page metadata.

Each loader returns a list of `PageText` (page_number is None for formats
that have no native pagination). Errors are raised as one of three typed
exceptions so callers can distinguish "bad input" from "bug":

- `UnsupportedFileTypeError`: extension isn't one we handle.
- `DocumentLoadError`: file exists and has a supported extension, but is
  corrupt/unreadable.
- `EmptyDocumentError`: file parsed fine but contains no extractable text.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentLoadError(Exception):
    """Raised when a file with a supported extension can't be parsed."""


class UnsupportedFileTypeError(Exception):
    """Raised when a file's extension isn't one of the supported types."""


class EmptyDocumentError(Exception):
    """Raised when a successfully-parsed document has no extractable text."""


@dataclass
class PageText:
    page_number: int | None
    text: str


@dataclass
class LoadedDocument:
    filename: str
    file_type: str
    pages: list[PageText]


def _load_pdf(path: Path) -> list[PageText]:
    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError, ValueError) as exc:
        raise DocumentLoadError(f"Could not read PDF '{path.name}': {exc}") from exc

    if reader.is_encrypted:
        raise DocumentLoadError(f"'{path.name}' is password-protected; cannot extract text.")

    pages: list[PageText] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pypdf can raise assorted errors on malformed page content
            raise DocumentLoadError(
                f"Could not extract text from page {index} of '{path.name}': {exc}"
            ) from exc
        pages.append(PageText(page_number=index, text=text))
    return pages


def _load_docx(path: Path) -> list[PageText]:
    try:
        doc = DocxDocument(str(path))
    except (PackageNotFoundError, OSError, ValueError) as exc:
        raise DocumentLoadError(f"Could not read DOCX '{path.name}': {exc}") from exc

    text = "\n\n".join(paragraph.text for paragraph in doc.paragraphs)
    return [PageText(page_number=None, text=text)]


def _load_txt(path: Path) -> list[PageText]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            text = path.read_text(encoding="latin-1")
        except OSError as exc:
            raise DocumentLoadError(f"Could not read text file '{path.name}': {exc}") from exc
    except OSError as exc:
        raise DocumentLoadError(f"Could not read text file '{path.name}': {exc}") from exc
    return [PageText(page_number=None, text=text)]


def _load_md(path: Path) -> list[PageText]:
    return _load_txt(path)


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
    ".md": _load_md,
}


def load_document(path: str | Path) -> LoadedDocument:
    """Load a single file and return its extracted text, grouped by page.

    Raises `UnsupportedFileTypeError`, `DocumentLoadError`, or
    `EmptyDocumentError` on failure; never raises a raw parser exception.
    """
    path = Path(path)
    if not path.exists() or not path.is_file():
        raise DocumentLoadError(f"File not found: '{path}'")

    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{suffix or '(none)'}' for '{path.name}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    pages = loader(path)

    if not pages or not any(page.text.strip() for page in pages):
        raise EmptyDocumentError(f"'{path.name}' contains no extractable text.")

    return LoadedDocument(filename=path.name, file_type=suffix.lstrip("."), pages=pages)
