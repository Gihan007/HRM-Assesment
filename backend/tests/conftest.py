"""Shared pytest fixtures.

Text fixture files (sample.txt, sample.md, corrupt.*, unsupported.xyz,
empty.txt) live as real files in tests/fixtures/ and are committed to
git. The PDF/DOCX fixtures are binary, so instead of committing binary
blobs they're generated on the fly into tests/fixtures/generated/
(gitignored) the first time a test asks for them.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

FIXTURES_DIR = Path(__file__).parent / "fixtures"
GENERATED_DIR = FIXTURES_DIR / "generated"

PDF_PAGE_TEXTS = [
    "Page one content. This document describes the RAG ingestion pipeline.",
    "Page two content. It continues the discussion from the previous page.",
    "Page three content. This is the final page of the sample document.",
]


@pytest.fixture(scope="session")
def fixtures_dir() -> Path:
    return FIXTURES_DIR


@pytest.fixture(scope="session")
def sample_pdf_path() -> Path:
    """A valid 3-page PDF with known, distinct text per page."""
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    path = GENERATED_DIR / "sample.pdf"

    c = canvas.Canvas(str(path), pagesize=letter)
    for page_text in PDF_PAGE_TEXTS:
        c.drawString(72, 720, page_text)
        c.showPage()
    c.save()
    return path


@pytest.fixture(scope="session")
def empty_pdf_path() -> Path:
    """A structurally valid PDF with a single blank page (no text)."""
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    path = GENERATED_DIR / "empty.pdf"

    c = canvas.Canvas(str(path), pagesize=letter)
    c.showPage()
    c.save()
    return path


@pytest.fixture(scope="session")
def sample_docx_path() -> Path:
    """A valid DOCX with a couple of paragraphs."""
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    path = GENERATED_DIR / "sample.docx"

    doc = DocxDocument()
    doc.add_paragraph("This is the first paragraph of the sample DOCX fixture.")
    doc.add_paragraph("This is the second paragraph, with different content.")
    doc.save(str(path))
    return path
